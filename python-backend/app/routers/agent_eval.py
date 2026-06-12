"""Agent eval router — benchmarks, datasets, test cases, runs.

Covers TS parity for all agentEvalProcedure endpoints.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel
from pydantic import Field as PField
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_db
from app.dependencies import get_current_user_id
from app.services.agent_eval.service import AgentEvalService

router = APIRouter(prefix="/api/agent-eval", tags=["Agent Eval"])


def _detect_dataset_format(pathname: str, filename: Optional[str], requested: Optional[str] = None) -> str:
    if requested and requested != "auto":
        return requested
    name = (filename or pathname).lower()
    if name.endswith(".jsonl"):
        return "jsonl"
    if name.endswith(".json"):
        return "json"
    if name.endswith(".csv"):
        return "csv"
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return "xlsx"
    return "csv"


def _read_dataset_file(pathname: str, fmt: str) -> str | bytes:
    path = Path(pathname)
    if not path.exists() or not path.is_file():
        raise HTTPException(400, f"Dataset file not found: {pathname}")
    return path.read_bytes() if fmt == "xlsx" else path.read_text(encoding="utf-8-sig")


def _parse_dataset_rows(content: str | bytes, *, fmt: str, preview: Optional[int] = None) -> dict[str, Any]:
    rows: list[dict[str, Any]]
    if fmt == "json":
        data = json.loads(content.decode("utf-8") if isinstance(content, bytes) else content)
        if isinstance(data, dict):
            data = data.get("data") or data.get("rows") or data.get("testCases") or data.get("test_cases") or []
        if not isinstance(data, list):
            raise HTTPException(400, "JSON dataset must be an array or contain a rows/data array")
        rows = [item if isinstance(item, dict) else {"input": item} for item in data]
    elif fmt == "jsonl":
        text = content.decode("utf-8") if isinstance(content, bytes) else content
        rows = [json.loads(line) for line in text.splitlines() if line.strip()]
    elif fmt == "xlsx":
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise HTTPException(400, f"XLSX parsing is unavailable: {exc}") from exc
        from io import BytesIO

        workbook = load_workbook(
            BytesIO(content if isinstance(content, bytes) else content.encode()),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(value or "") for value in (values[0] if values else [])]
        rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        text = content.decode("utf-8") if isinstance(content, bytes) else content
        rows = list(csv.DictReader(text.splitlines()))

    headers = list(rows[0].keys()) if rows else []
    visible_rows = rows[:preview] if preview is not None else rows
    return {"headers": headers, "rows": visible_rows, "totalCount": len(rows), "format": fmt}


# ── Schemas ──────────────────────────────────────────────────────────

class BenchmarkCreate(BaseModel):
    model_config = {"populate_by_name": True}
    identifier: Optional[str] = None
    name: str
    description: Optional[str] = None
    metadata: Optional[dict[str, Any]] = None
    tags: Optional[list[str]] = None
    rubrics: Optional[list[Any]] = None
    # Legacy field
    agent_id: Optional[str] = PField(default=None, alias="agentId")
    config: Optional[dict[str, Any]] = None


class DatasetCreate(BaseModel):
    model_config = {"populate_by_name": True}
    benchmark_id: str = PField(alias="benchmarkId")
    name: str
    identifier: Optional[str] = None
    description: Optional[str] = None
    metadata: Optional[dict[str, Any]] = None
    eval_mode: Optional[str] = PField(default=None, alias="evalMode")
    eval_config: Optional[dict[str, Any]] = PField(default=None, alias="evalConfig")


class TestCaseCreate(BaseModel):
    model_config = {"populate_by_name": True}
    dataset_id: str = PField(alias="datasetId")
    content: Optional[dict[str, Any]] = None
    eval_mode: Optional[str] = PField(default=None, alias="evalMode")
    eval_config: Optional[dict[str, Any]] = PField(default=None, alias="evalConfig")
    metadata: Optional[dict[str, Any]] = None
    # Legacy fields
    input: Optional[str] = None
    expected_output: Optional[str] = None


class RunCreate(BaseModel):
    model_config = {"populate_by_name": True}
    dataset_id: Optional[str] = PField(default=None, alias="datasetId")
    name: Optional[str] = None
    target_agent_id: Optional[str] = PField(default=None, alias="targetAgentId")
    config: Optional[dict[str, Any]] = None
    # Legacy field
    benchmark_id: Optional[str] = PField(default=None, alias="benchmarkId")


# ── Benchmark endpoints ──────────────────────────────────────────────

@router.get("/benchmarks")
async def list_benchmarks(
    agent_id: Optional[str] = None,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    benchmarks = await svc.list_benchmarks(agent_id)
    return [
        {
            "id": b.id,
            "agent_id": b.agent_id,
            "name": b.name,
            "description": b.description,
            "created_at": b.created_at.isoformat() if b.created_at else None,
        }
        for b in benchmarks
    ]


@router.post("/benchmarks")
async def create_benchmark(
    body: BenchmarkCreate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    bench = await svc.create_benchmark(**body.model_dump(exclude_none=True))
    await session.commit()
    return {"id": bench.id}


@router.delete("/benchmarks/{benchmark_id}")
async def delete_benchmark(
    benchmark_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    await svc.delete_benchmark(benchmark_id)
    await session.commit()
    return {"ok": True}


# ── Dataset endpoints ────────────────────────────────────────────────

@router.get("/datasets")
async def list_datasets(
    benchmark_id: str = Query(default=None, alias="benchmarkId"),
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    datasets = await svc.list_datasets(benchmark_id)
    return [
        {
            "id": d.id,
            "benchmark_id": d.benchmark_id,
            "name": d.name,
            "description": d.description,
            "created_at": d.created_at.isoformat() if d.created_at else None,
        }
        for d in datasets
    ]


@router.post("/datasets")
async def create_dataset(
    body: DatasetCreate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    ds = await svc.create_dataset(**body.model_dump(exclude_none=True))
    await session.commit()
    return {"id": ds.id}


# ── Test Case endpoints ──────────────────────────────────────────────

@router.get("/test-cases")
async def list_test_cases(
    dataset_id: str = Query(default=None, alias="datasetId"),
    limit: int = Query(default=50, le=200),
    offset: int = Query(default=0),
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    cases = await svc.list_test_cases(dataset_id, limit=limit, offset=offset)
    return [
        {
            "id": c.id,
            "input": c.input,
            "expected_output": c.expected_output,
            "content": getattr(c, "content", None),
            "metadata": getattr(c, "metadata_", None),
            "sortOrder": getattr(c, "sort_order", None),
            "created_at": c.created_at.isoformat() if c.created_at else None,
        }
        for c in cases
    ]


@router.post("/test-cases")
async def create_test_case(
    body: TestCaseCreate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    input_text = body.input
    expected = body.expected_output
    if body.content:
        input_text = input_text or body.content.get("input")
        expected = expected or body.content.get("expected")
    tc = await svc.create_test_case(
        dataset_id=body.dataset_id,
        input=input_text,
        expected_output=expected,
        metadata_=body.metadata,
    )
    await session.commit()
    return {"id": tc.id}


@router.post("/test-cases/batch")
async def batch_create_test_cases(
    dataset_id: str = Query(...),
    body: list[TestCaseCreate] = ...,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    cases = await svc.batch_create_test_cases(
        dataset_id,
        [c.model_dump() for c in body],
    )
    await session.commit()
    return {"created": len(cases)}


# ── Run endpoints ────────────────────────────────────────────────────

@router.get("/runs")
async def list_runs(
    benchmark_id: Optional[str] = Query(default=None, alias="benchmarkId"),
    dataset_id: Optional[str] = Query(default=None, alias="datasetId"),
    limit: int = Query(default=50, le=200),
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    runs = await svc.list_runs(benchmark_id, limit=limit, dataset_id=dataset_id)
    return [
        {
            "id": r.id,
            "benchmark_id": r.benchmark_id,
            "dataset_id": r.dataset_id,
            "status": r.status,
            "results": r.results,
            "error": r.error,
            "started_at": r.started_at.isoformat() if r.started_at else None,
            "completed_at": r.completed_at.isoformat() if r.completed_at else None,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in runs
    ]


@router.post("/runs")
async def create_run(
    body: RunCreate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    run = await svc.create_run(**body.model_dump(exclude_none=True))
    await session.commit()
    return {"id": run.id}


@router.get("/runs/{run_id}")
async def get_run(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status == "running":
        await svc.check_and_handle_run_timeout(run_id)
        await session.commit()
        run = await svc.get_run(run_id)
        if not run:
            raise HTTPException(404, "Run not found")
    return {
        "id": run.id,
        "benchmark_id": run.benchmark_id,
        "status": run.status,
        "results": run.results,
        "error": run.error,
        "created_at": run.created_at.isoformat() if run.created_at else None,
    }


@router.post("/runs/{run_id}/execute")
async def execute_run(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    try:
        result = await svc.execute_run(run_id)
        await session.commit()
        return result
    except ValueError as exc:
        raise HTTPException(404, str(exc))


@router.get("/runs/{run_id}/topics")
async def list_run_topics(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    topics = await svc.list_run_topics(run_id)
    return [
        {
            "id": t.id,
            "test_case_id": t.test_case_id,
            "status": t.status,
            "score": t.score,
            "created_at": t.created_at.isoformat() if t.created_at else None,
        }
        for t in topics
    ]


# ── Missing TS parity endpoints ────────────────────────────────────


class BenchmarkUpdate(BaseModel):
    model_config = {"populate_by_name": True}
    identifier: Optional[str] = None
    name: Optional[str] = None
    description: Optional[str] = None
    metadata: Optional[dict[str, Any]] = None
    tags: Optional[list[str]] = None
    config: Optional[dict[str, Any]] = None


class DatasetUpdate(BaseModel):
    model_config = {"populate_by_name": True}
    name: Optional[str] = None
    description: Optional[str] = None
    metadata: Optional[dict[str, Any]] = None
    eval_mode: Optional[str] = PField(default=None, alias="evalMode")
    eval_config: Optional[dict[str, Any]] = PField(default=None, alias="evalConfig")


class TestCaseUpdate(BaseModel):
    model_config = {"populate_by_name": True}
    content: Optional[dict[str, Any]] = None
    metadata: Optional[dict[str, Any]] = None
    eval_mode: Optional[str] = PField(default=None, alias="evalMode")
    eval_config: Optional[dict[str, Any]] = PField(default=None, alias="evalConfig")
    sort_order: Optional[int] = PField(default=None, alias="sortOrder")


class RunUpdate(BaseModel):
    model_config = {"populate_by_name": True}
    name: Optional[str] = None
    dataset_id: Optional[str] = PField(default=None, alias="datasetId")
    target_agent_id: Optional[str] = PField(default=None, alias="targetAgentId")
    config: Optional[dict[str, Any]] = None


class UpdateRunStatusBody(BaseModel):
    status: str
    error: Optional[str] = None


class UpdateRunMetricsBody(BaseModel):
    metrics: dict[str, Any]


class ParseDatasetFileBody(BaseModel):
    model_config = {"populate_by_name": True}
    pathname: str
    filename: Optional[str] = None
    format: Optional[str] = None


class ImportDatasetBody(BaseModel):
    model_config = {"populate_by_name": True}
    dataset_id: str = PField(alias="datasetId")
    pathname: str
    filename: Optional[str] = None
    format: Optional[str] = None
    field_mapping: Optional[dict[str, Any]] = PField(default=None, alias="fieldMapping")
    # Legacy
    test_cases: Optional[list[dict[str, Any]]] = PField(default=None, alias="testCases")


class RetryRunCaseBody(BaseModel):
    model_config = {"populate_by_name": True}
    test_case_id: str = PField(alias="testCaseId")


class ResumeRunCaseBody(BaseModel):
    model_config = {"populate_by_name": True}
    test_case_id: str = PField(alias="testCaseId")
    thread_id: Optional[str] = PField(default=None, alias="threadId")


class BatchResumeBody(BaseModel):
    targets: list[dict[str, Any]]


# ── Benchmark: get, update ──────────────────────────────────────────


@router.get("/benchmarks/{benchmark_id}")
async def get_benchmark(
    benchmark_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    bench = await svc.get_benchmark(benchmark_id)
    if not bench:
        raise HTTPException(404, "Benchmark not found")
    return {
        "id": bench.id,
        "agent_id": bench.agent_id,
        "name": bench.name,
        "description": bench.description,
        "config": bench.config,
        "created_at": bench.created_at.isoformat() if bench.created_at else None,
    }


@router.patch("/benchmarks/{benchmark_id}")
@router.put("/benchmarks/{benchmark_id}")
async def update_benchmark(
    benchmark_id: str,
    body: BenchmarkUpdate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    result = await svc.update_benchmark(benchmark_id, **body.model_dump(exclude_none=True))
    await session.commit()
    if not result:
        raise HTTPException(404, "Benchmark not found")
    return {"ok": True}


# ── Dataset: get, update, delete, parseFile, import ─────────────────


@router.get("/datasets/{dataset_id}")
async def get_dataset(
    dataset_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    ds = await svc.get_dataset(dataset_id)
    if not ds:
        raise HTTPException(404, "Dataset not found")
    return {
        "id": ds.id,
        "benchmark_id": ds.benchmark_id,
        "name": ds.name,
        "description": ds.description,
        "created_at": ds.created_at.isoformat() if ds.created_at else None,
    }


@router.patch("/datasets/{dataset_id}")
@router.put("/datasets/{dataset_id}")
async def update_dataset(
    dataset_id: str,
    body: DatasetUpdate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    result = await svc.update_dataset(dataset_id, **body.model_dump(exclude_none=True))
    await session.commit()
    if not result:
        raise HTTPException(404, "Dataset not found")
    return {"ok": True}


@router.delete("/datasets/{dataset_id}")
async def delete_dataset(
    dataset_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    await svc.delete_dataset(dataset_id)
    await session.commit()
    return {"ok": True}


@router.post("/datasets/parse-file")
async def parse_dataset_file(
    body: ParseDatasetFileBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Parse a local dataset file and return preview rows."""
    fmt = _detect_dataset_format(body.pathname, body.filename, body.format)
    content = _read_dataset_file(body.pathname, fmt)
    parsed = _parse_dataset_rows(content, fmt=fmt, preview=50)
    return {
        "headers": parsed["headers"],
        "preview": parsed["rows"],
        "totalCount": parsed["totalCount"],
        "format": parsed["format"],
    }


@router.post("/datasets/import")
async def import_dataset(
    body: ImportDatasetBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Batch import test cases into a dataset."""
    svc = AgentEvalService(session, user_id)
    if body.test_cases is not None:
        case_inputs = body.test_cases
    else:
        fmt = _detect_dataset_format(body.pathname, body.filename, body.format)
        content = _read_dataset_file(body.pathname, fmt)
        parsed = _parse_dataset_rows(content, fmt=fmt)
        mapping = body.field_mapping or {}
        input_col = mapping.get("input")
        if not input_col:
            raise HTTPException(400, "fieldMapping.input is required")
        expected_col = mapping.get("expected")
        expected_delimiter = mapping.get("expectedDelimiter")
        case_inputs = []
        for row in parsed["rows"]:
            expected_output = None
            if expected_col and row.get(expected_col) is not None:
                expected_raw = str(row.get(expected_col))
                if expected_delimiter:
                    parts = [part.strip() for part in expected_raw.split(expected_delimiter) if part.strip()]
                    expected_output = json.dumps(parts) if len(parts) > 1 else expected_raw
                else:
                    expected_output = expected_raw
            metadata = {}
            if isinstance(mapping.get("metadata"), dict):
                metadata = {key: row.get(column) for key, column in mapping["metadata"].items()}
            case_inputs.append(
                {
                    "input": str(row.get(input_col) or ""),
                    "expected_output": expected_output,
                    "metadata": metadata,
                }
            )
    cases = await svc.batch_create_test_cases(body.dataset_id, case_inputs)
    await session.commit()
    return {"count": len(cases), "data": cases}


# ── TestCase: get, update, delete ───────────────────────────────────


@router.get("/test-cases/{test_case_id}")
async def get_test_case(
    test_case_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    tc = await svc.get_test_case(test_case_id)
    if not tc:
        raise HTTPException(404, "Test case not found")
    return {
        "id": tc.id,
        "input": tc.input,
        "expected_output": tc.expected_output,
        "metadata": getattr(tc, "metadata_", None),
        "created_at": tc.created_at.isoformat() if tc.created_at else None,
    }


@router.patch("/test-cases/{test_case_id}")
@router.put("/test-cases/{test_case_id}")
async def update_test_case(
    test_case_id: str,
    body: TestCaseUpdate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    result = await svc.update_test_case(test_case_id, **body.model_dump(exclude_none=True))
    await session.commit()
    if not result:
        raise HTTPException(404, "Test case not found")
    return {"ok": True}


@router.delete("/test-cases/{test_case_id}")
async def delete_test_case(
    test_case_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    await svc.delete_test_case(test_case_id)
    await session.commit()
    return {"ok": True}


# ── Run: getDetails, delete, start, abort, retry, resume, progress ──


@router.get("/runs/{run_id}/details")
async def get_run_details(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    result = await svc.get_run_details(run_id)
    if not result:
        raise HTTPException(404, "Run not found")
    return result


@router.delete("/runs/{run_id}")
async def delete_run(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    svc = AgentEvalService(session, user_id)
    await svc.delete_run(run_id)
    await session.commit()
    return {"ok": True}


class StartRunBody(BaseModel):
    force: Optional[bool] = None


@router.post("/runs/{run_id}/start")
async def start_run(
    run_id: str,
    body: Optional[StartRunBody] = None,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Start executing an eval run (transitions idle/failed → pending → running)."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status not in ("idle", "failed"):
        raise HTTPException(400, f"Run cannot be started from status: {run.status}")
    await svc.update_run_status(run_id, "pending")
    try:
        result = await svc.execute_run(run_id)
        await session.commit()
        return result
    except Exception as exc:
        await svc.update_run_status(run_id, "failed", error=str(exc))
        await session.commit()
        raise HTTPException(500, str(exc))


@router.post("/runs/{run_id}/abort")
async def abort_run(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Abort a running evaluation."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status not in ("running", "pending"):
        raise HTTPException(400, f"Run is not running (status: {run.status})")
    await svc.update_run_status(run_id, "canceled")
    await session.commit()
    return {"success": True}


@router.post("/runs/{run_id}/retry-errors")
async def retry_run_errors(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Retry all failed test cases in a run."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    retry_count = await svc.retry_failed_cases(run_id)
    await session.commit()
    return {"success": True, "retryCount": retry_count, "runId": run_id}


@router.post("/runs/{run_id}/retry-case")
async def retry_run_case(
    run_id: str,
    body: RetryRunCaseBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Retry a specific test case in a run."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    await svc.retry_case(run_id, body.test_case_id)
    await session.commit()
    return {"success": True, "runId": run_id, "testCaseId": body.test_case_id}


@router.post("/runs/{run_id}/resume-case")
async def resume_run_case(
    run_id: str,
    body: ResumeRunCaseBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Resume a specific test case in a run."""
    svc = AgentEvalService(session, user_id)
    result = await svc.resume_case(run_id, body.test_case_id, thread_id=body.thread_id)
    await session.commit()
    return result


@router.post("/runs/{run_id}/batch-resume")
async def batch_resume_run_cases(
    run_id: str,
    body: BatchResumeBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Batch resume multiple test cases."""
    svc = AgentEvalService(session, user_id)
    succeeded = 0
    failed = 0
    for target in body.targets:
        try:
            await svc.resume_case(run_id, target["testCaseId"], thread_id=target.get("threadId"))
            succeeded += 1
        except Exception:
            failed += 1
    await session.commit()
    return {"succeeded": succeeded, "failed": failed, "total": len(body.targets)}


@router.get("/runs/{run_id}/resumable-cases")
async def get_resumable_cases(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Get test cases that can be resumed."""
    svc = AgentEvalService(session, user_id)
    cases = await svc.get_resumable_cases(run_id)
    return cases


@router.get("/runs/{run_id}/progress")
async def get_run_progress(
    run_id: str,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Get real-time progress of a running evaluation."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    topics = await svc.list_run_topics(run_id)
    total = len(topics)
    completed = sum(1 for t in topics if t.status in ("completed", "failed", "canceled"))
    running = sum(1 for t in topics if t.status == "running")
    pending = sum(1 for t in topics if t.status == "pending")
    failed = sum(1 for t in topics if t.status == "failed")
    return {
        "status": run.status,
        "total": total,
        "completed": completed,
        "running": running,
        "pending": pending,
        "failed": failed,
        "progress": (completed / total * 100) if total > 0 else 0,
    }


@router.get("/runs/{run_id}/results")
async def get_run_results(
    run_id: str,
    limit: int = Query(default=50, le=200),
    offset: int = Query(default=0),
    status_filter: Optional[str] = Query(default=None, alias="status"),
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Get detailed results of test case executions."""
    svc = AgentEvalService(session, user_id)
    results = await svc.get_run_results(run_id, limit=limit, offset=offset, status_filter=status_filter)
    return results


# ── Run: updateStatus, updateMetrics, update ────────────────────────


@router.put("/runs/{run_id}/status")
async def update_run_status(
    run_id: str,
    body: UpdateRunStatusBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Update run status (internal use)."""
    svc = AgentEvalService(session, user_id)
    await svc.update_run_status(run_id, body.status, error=body.error)
    await session.commit()
    return {"ok": True}


@router.put("/runs/{run_id}/metrics")
async def update_run_metrics(
    run_id: str,
    body: UpdateRunMetricsBody,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Update run metrics (internal use)."""
    svc = AgentEvalService(session, user_id)
    await svc.update_run_metrics(run_id, body.metrics)
    await session.commit()
    return {"ok": True}


@router.patch("/runs/{run_id}")
@router.put("/runs/{run_id}")
async def update_run(
    run_id: str,
    body: RunUpdate,
    user_id: str = Depends(get_current_user_id),
    session: AsyncSession = Depends(get_db),
):
    """Update run (user-facing: name, datasetId, targetAgentId, config)."""
    svc = AgentEvalService(session, user_id)
    run = await svc.get_run(run_id)
    if not run:
        raise HTTPException(404, "Run not found")
    if run.status not in ("idle", "failed") and (body.dataset_id or body.target_agent_id):
        raise HTTPException(400, "Cannot change datasetId or targetAgentId after run has started")
    result = await svc.update_run(run_id, **body.model_dump(exclude_none=True))
    await session.commit()
    if not result:
        raise HTTPException(404, "Run not found")
    return {"ok": True}
